#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
备考数据中枢 —— 2027 考研管综(199) + 英语二 + CET-6
====================================================

一个脚本搞定：错题归因、时长统计、模考追踪、知识点覆盖、遗忘曲线复习提醒、
Excel 仪表盘、每日分钟级时间表。

用法速查
--------
    python tools/tracker.py init                 # 初始化数据文件
    python tools/tracker.py today                # 打印今日时间表（自动选版本）
                                                 #   周一~周五 → B(有课型)
                                                 #   周六      → A(全天型)
                                                 #   周日      → C(复盘+六级听力)
    python tools/tracker.py today -d A           # 强制指定版本
    python tools/tracker.py seed-baseline        # 导入 2017-2021 真题错点基线
    python tools/tracker.py plan --export        # 导出时间表到 plans/每日时间表.md
    python tools/tracker.py add-mistake          # 交互式录入错题
    python tools/tracker.py add-session          # 交互式录入学习时长
    python tools/tracker.py add-paper            # 交互式录入模考成绩
    python tools/tracker.py review               # 查看今日到期复习的错题
    python tools/tracker.py review --done 3      # 把 3 条到期错题标记为已复习
    python tools/tracker.py report               # 控制台报表
    python tools/tracker.py excel                # 生成 report/备考仪表盘.xlsx

数据全部落在 data/*.csv，可直接用 Excel 打开手改。
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

# --------------------------------------------------------------------------- #
# 路径与常量
# --------------------------------------------------------------------------- #
try:  # Windows 控制台中文输出
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
REPORT = ROOT / "report"
PLANS = ROOT / "plans"

CONFIG_PATH = DATA / "config.json"

DEFAULT_CONFIG = {
    "exam_kaoyan": "2026-12-19",
    "exam_cet6_written": "2026-12-12",
    "exam_cet6_oral": "2026-11-22",
    "targets": {
        "数学": 63, "逻辑": 52, "写作": 45,
        "管综总分": 160, "英语二": 78, "六级": 550,
    },
}

# 科目 → 模块映射
MODULES = {
    "数学": ["算术", "代数", "几何", "数据分析", "条件充分性判断"],
    "逻辑": ["形式逻辑", "分析推理", "论证逻辑"],
    "写作": ["论证有效性分析", "论说文"],
    "英语二": ["完形", "阅读", "新题型", "翻译", "写作"],
    "六级": ["听力", "阅读", "写作", "翻译", "词汇"],
}

ERROR_TYPES = [
    "概念不清", "方法缺失", "计算失误", "审题偏差",
    "陷阱误判", "时间不足", "表达不规范",
]

# 艾宾浩斯复习间隔（天）
REVIEW_INTERVALS = [1, 2, 4, 7, 15, 30]

# CSV 表头（中文，方便直接用 Excel 打开）
HEADERS = {
    "mistakes": [
        "错题ID", "录入日期", "科目", "模块", "考点", "来源", "错误类型",
        "根本原因", "避坑口诀", "复习阶段", "下次复习", "已清零", "备注",
    ],
    "sessions": [
        "日期", "科目", "内容", "实际时长(分)", "计划时长(分)", "专注度(1-5)", "备注",
    ],
    "papers": [
        "日期", "科目", "试卷名称", "用时(分)", "得分", "满分", "错题数", "备注",
    ],
    "coverage": [
        "科目", "模块", "考点", "掌握度(1-5)", "最近练习", "累计错次", "备注",
    ],
}

# --------------------------------------------------------------------------- #
# 每日时间表（分钟级）—— 单一数据源，`tracker.py plan --export` 由此生成 markdown
# --------------------------------------------------------------------------- #
SCHEDULE = {
    "A": {
        "title": "版本 A｜全天可支配日（周末 / 无课日 / 冲刺期）",
        "total": "净学习约 10.5 小时",
        "rows": [
            ("06:20-06:35", "起床 · 洗漱 · 温水", "不碰手机，直接下床", "必修"),
            ("06:35-06:50", "昨日错题卡速览", "只看「避坑口诀」列，3 分钟一张卡", "必修"),
            ("06:50-07:35", "英语词汇 45min", "新词 50 + 复习 100，六级/考研双轨词表", "必修"),
            ("07:35-08:05", "早餐", "可放英语泛听（六级真题音频，不求听懂）", "必修"),
            ("08:05-08:15", "数学热身 10min", "公式默写 + 速算 10 题（几何公式/均值不等式）", "必修"),
            ("08:15-09:45", "数学专项 90min", "单日几何(截面/割补) · 双日概率与排列组合", "核心"),
            ("09:45-10:00", "休息 15min", "起身活动，不看屏幕", "必修"),
            ("10:00-11:30", "逻辑专项 90min", "形式逻辑 → 分析推理 → 论证逻辑 三日轮转", "核心"),
            ("11:30-11:45", "上午复盘 15min", "口述今日两个新学的解题入口", "必修"),
            ("11:45-12:40", "午餐 + 散步", "饭后走动 10 分钟，不立刻坐下", "必修"),
            ("12:40-13:20", "午休 40min", "定闹钟，绝不超时", "必修"),
            ("13:20-13:30", "醒神", "冷水洗脸 / 拉伸，写一句今日目标", "必修"),
            ("13:30-15:10", "英语二精读 100min", "2 篇真题阅读：限时 → 对答案 → 逐句拆长难句", "核心"),
            ("15:10-15:25", "休息 15min", "远眺，护眼", "必修"),
            ("15:25-16:25", "六级听力专项 60min", "精听：盲听→逐句复听→对原文标生词→跟读", "核心"),
            ("16:25-16:40", "休息 15min", "补充能量", "必修"),
            ("16:40-18:10", "管综写作 90min", "单日论证有效性分析（限时 30min 成文）· 双日论说文", "核心"),
            ("18:10-19:10", "晚餐 + 完全放松", "不学习，恢复精力", "必修"),
            ("19:10-20:30", "错题重做 80min", "只做已标注的错题，不看答案先独立推导", "核心"),
            ("20:30-21:10", "英语写作/翻译 40min", "小作文背框架 · 大作文图表段；英二翻译逐句对照", "核心"),
            ("21:10-21:45", "当日复盘 + 数据录入", "跑 tracker.py 录错题与时长，写明日 3 个重点", "必修"),
            ("21:45-22:15", "睡前词汇/泛听", "当天新词回滚一遍，音频不设任务量", "必修"),
            ("22:15-22:40", "洗漱 · 睡前准备", "手机放远离床头", "必修"),
            ("22:40", "熄灯", "保证 7.5 小时睡眠 —— 睡眠是记忆固化环节", "铁律"),
        ],
    },
    "B": {
        "title": "版本 B｜有课 / 实习日（净学习约 6 小时）",
        "total": "净学习约 6 小时，靠碎片时间补",
        "rows": [
            ("06:30-07:15", "英语词汇 45min", "新词 30 + 复习 80（有课日减量不减频）", "必修"),
            ("07:15-07:50", "早餐 + 通勤", "耳机放六级听力，只求语感不设任务", "碎片"),
            ("07:50-12:00", "课程", "课间 5min ×4：翻词汇卡 / 做 1 道逻辑小题", "碎片"),
            ("12:00-13:00", "午餐 + 午休 25min", "必须小睡，下午效率靠它", "必修"),
            ("13:00-14:00", "数学专项 60min", "只攻一个模块，宁可少而透", "核心"),
            ("14:00-18:00", "课程", "课间继续碎片：英语长难句 1 句", "碎片"),
            ("18:00-19:00", "晚餐 + 放松", "完全不学习", "必修"),
            ("19:00-20:30", "英语精读 + 听力 90min", "阅读 1 篇精读 60min + 听力精听 30min", "核心"),
            ("20:30-21:30", "逻辑 / 写作轮转 60min", "单日逻辑 · 双日写作", "核心"),
            ("21:30-22:00", "错题复盘 + 数据录入", "tracker.py 录入，别攒到周末", "必修"),
            ("22:00-22:30", "词汇复习", "当天新词回滚", "必修"),
            ("22:30-23:00", "洗漱 · 熄灯", "23:00 前必须睡", "铁律"),
        ],
    },
    "sunday": {
        "title": "版本 C｜周日｜复盘与六级听力专项日",
        "total": "不排新内容，只做清零与检测",
        "rows": [
            ("07:00-07:45", "英语词汇", "本周生词本全量回滚", "必修"),
            ("07:45-08:30", "早餐 + 到自习室", "—", "必修"),
            ("08:30-10:00", "六级听力专项精听 90min", "1 套真题：盲听 → 逐句复听 → 原文标生词 → 跟读 3 遍", "核心"),
            ("10:00-10:15", "休息", "—", "必修"),
            ("10:15-11:45", "数学「3+3+1」的「1」：错题重做清零 90min", "本周错题横向重做，做对才准划掉", "核心"),
            ("11:45-13:30", "午餐 + 午休", "—", "必修"),
            ("13:30-15:30", "管综全科小模考 120min", "数学 60min + 逻辑 60min，严格计时", "核心"),
            ("15:30-16:00", "休息", "—", "必修"),
            ("16:00-17:30", "模考精析 90min", "错因归档 + 写「一题三改」", "核心"),
            ("17:30-18:30", "晚餐 + 放松", "—", "必修"),
            ("18:30-20:00", "英语二真题 1 套（阅读+新题型）90min", "限时做，当周进度验收", "核心"),
            ("20:00-21:00", "写作输出 60min", "本周 1 篇论证有效性分析 + 1 篇论说文交替", "核心"),
            ("21:00-21:40", "周复盘 + 下周计划", "跑 tracker.py report，看薄弱模块，定下周 3 个攻坚点", "必修"),
            ("21:40-22:30", "自由 / 休息", "给自己留一段完全放空的时间", "必修"),
        ],
    },
}

STRATEGY = {
    "数学": "3+3+1 周节奏：3 套限时模考 + 3 次深度精析（写明错因与几何/概率突破口）+ 1 次错题重做清零",
    "逻辑": "零失误：形式逻辑用假言逆否(A→B ≡ ¬B→¬A)与假言选言转换(A→B ≡ ¬A∨B)；分析推理找最大频次项/确定性信息；论证逻辑画「论据→论点」链条",
    "英语": "一套英语两个出口：阅读用英二真题精读，听力只有六级要练，写作共用素材库",
    "写作": "论证有效性分析：标题+开头+4 段正文（漏洞+原因+推不出）+结尾；论说文：5 分钟审题立意 + 引-议-联-结",
}

# 2017–2021 年真题错题分析反推出的已知错点（(科目, 模块, 考点)）
# 用途：在还没有模考分数时，先把「知识点覆盖表」填上待校准的骨架。
BASELINE_COVERAGE = [
    ("数学", "几何", "正六边形阴影面积（割补 / 整体减空白）"),
    ("数学", "几何", "球内接正方体与大圆截面"),
    ("数学", "几何", "圆内接三角形"),
    ("数学", "几何", "三角形中线长公式"),
    ("数学", "几何", "解析几何不等式与圆的切线"),
    ("数学", "数据分析", "集合容斥原理"),
    ("数学", "数据分析", "动态路径计数（少步数）"),
    ("数学", "数据分析", "分组问题（均分需除以 A_k^k）"),
    ("数学", "数据分析", "概率不等式的充分性判断"),
    ("数学", "条件充分性判断", "单条件充分 vs 联合充分（C 选项）边界"),
    ("数学", "代数", "平均数 / 平均分变化"),
    ("数学", "代数", "均值不等式及等号成立条件"),
    ("数学", "代数", "不定方程应用"),
    ("数学", "代数", "数列 S_n 与 a_n 关系（n=1 边界）"),
    ("数学", "代数", "分段绝对值函数最小值"),
    ("数学", "算术", "植树问题中的角上重复计数"),
    ("逻辑", "形式逻辑", "假言命题逆否等价 A→B ≡ ¬B→¬A"),
    ("逻辑", "形式逻辑", "充分/必要条件转换"),
    ("逻辑", "分析推理", "多条件约束的标准化画板 / 列表法"),
    ("逻辑", "分析推理", "最大频次项 / 确定性信息作为突破口"),
    ("逻辑", "论证逻辑", "比例数量陷阱（绝对量 vs 相对量）"),
    ("逻辑", "论证逻辑", "因果倒置识别"),
    ("逻辑", "论证逻辑", "偷换概念识别"),
    ("逻辑", "论证逻辑", "搭桥法补前提"),
]


# --------------------------------------------------------------------------- #
# 通用 IO
# --------------------------------------------------------------------------- #
def load_config() -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))  # 深拷贝
    if CONFIG_PATH.exists():
        try:
            user = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            cfg.update({k: v for k, v in user.items() if k != "targets"})
            cfg["targets"].update(user.get("targets", {}))
        except Exception as exc:
            print(f"[警告] config.json 解析失败，使用默认配置：{exc}")
    return cfg


def ensure_dirs() -> None:
    for p in (DATA, REPORT, PLANS):
        p.mkdir(parents=True, exist_ok=True)


def csv_path(name: str) -> Path:
    return DATA / f"{name}.csv"


def init_files(force: bool = False) -> None:
    ensure_dirs()
    for name, header in HEADERS.items():
        path = csv_path(name)
        if path.exists() and not force:
            print(f"[跳过] {path.name} 已存在")
            continue
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerow(header)
        print(f"[创建] {path.name}")
    if not CONFIG_PATH.exists() or force:
        CONFIG_PATH.write_text(
            json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print("[创建] config.json")


def read_rows(name: str) -> list[dict]:
    path = csv_path(name)
    if not path.exists():
        return []
    with path.open("r", newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.DictReader(fh) if any((v or "").strip() for v in row.values())]


def append_row(name: str, row: dict) -> None:
    init_files()
    path = csv_path(name)
    with path.open("a", newline="", encoding="utf-8-sig") as fh:
        csv.writer(fh).writerow([row.get(h, "") for h in HEADERS[name]])


def replace_rows(name: str, rows: list[dict]) -> None:
    path = csv_path(name)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(HEADERS[name])
        for row in rows:
            writer.writerow([row.get(h, "") for h in HEADERS[name]])


def ask(prompt: str, default: str = "", choices: list[str] | None = None) -> str:
    """交互式询问；choices 提供编号选项。"""
    if choices:
        print(f"  {prompt}")
        for i, c in enumerate(choices, 1):
            print(f"    {i}. {c}")
        prompt = "  请输入编号"
    suffix = f" [{default}]" if default else ""
    while True:
        raw = input(f"{prompt}{suffix}: ").strip()
        if not raw:
            if default:
                return default
            print("  ⚠ 不能为空")
            continue
        if choices:
            if raw in choices:
                return raw
            if raw.isdigit() and 1 <= int(raw) <= len(choices):
                return choices[int(raw) - 1]
            print("  ⚠ 无效编号")
            continue
        return raw


def to_int(text: str, default: int = 0) -> int:
    try:
        return int(float(str(text).strip()))
    except Exception:
        return default


def to_float(text: str, default: float = 0.0) -> float:
    try:
        return float(str(text).strip())
    except Exception:
        return default


def parse_date(text: str) -> date:
    text = (text or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m-%d", "%m/%d"):
        try:
            parsed = datetime.strptime(text, fmt).date()
            if fmt.startswith("%m"):
                parsed = parsed.replace(year=date.today().year)
            return parsed
        except ValueError:
            continue
    return date.today()


# --------------------------------------------------------------------------- #
# 错题
# --------------------------------------------------------------------------- #
def cmd_add_mistake(args) -> None:
    rows = read_rows("mistakes")
    if not args.yes:
        print("\n=== 录入错题（直接回车 = 用默认值 / 跳过）===")
        subject = ask("科目", choices=list(MODULES.keys()))
        module = ask("模块", choices=MODULES[subject])
        point = ask("考点（如：球内接正方体截面）")
        source = ask("来源（如：2021 真题第 12 题）")
        etype = ask("错误类型", choices=ERROR_TYPES)
        cause = ask("根本原因（我为什么错）")
        fix = ask("避坑口诀（下次怎么做）")
        note = ask("备注", default="-")
    else:
        subject, module, point, source = args.subject, args.module, args.point, args.source
        etype, cause, fix, note = args.error_type, args.cause, args.fix, args.note or "-"

    today = date.today()
    row = {
        "错题ID": f"E{len(rows) + 1:04d}",
        "录入日期": today.isoformat(),
        "科目": subject, "模块": module, "考点": point, "来源": source,
        "错误类型": etype, "根本原因": cause, "避坑口诀": fix,
        "复习阶段": 0,
        "下次复习": (today + timedelta(days=REVIEW_INTERVALS[0])).isoformat(),
        "已清零": "否", "备注": note,
    }
    append_row("mistakes", row)
    print(f"\n[已录入] {row['错题ID']} | {subject}·{module}·{point}")
    print(f"           下次复习：{row['下次复习']}（首次间隔 1 天）")

    # 同步知识点覆盖表
    sync_coverage(subject, module, point)


def sync_coverage(subject: str, module: str, point: str) -> None:
    rows = read_rows("coverage")
    for row in rows:
        if row["科目"] == subject and row["模块"] == module and row["考点"] == point:
            row["累计错次"] = str(to_int(row.get("累计错次", "0")) + 1)
            row["最近练习"] = date.today().isoformat()
            replace_rows("coverage", rows)
            return
    rows.append({
        "科目": subject, "模块": module, "考点": point,
        "掌握度(1-5)": "2", "最近练习": date.today().isoformat(),
        "累计错次": "1", "备注": "",
    })
    replace_rows("coverage", rows)


def due_mistakes(today: date | None = None) -> list[dict]:
    today = today or date.today()
    out = []
    for row in read_rows("mistakes"):
        if (row.get("已清零") or "否") == "是":
            continue
        try:
            nxt = datetime.strptime(row.get("下次复习", ""), "%Y-%m-%d").date()
        except ValueError:
            continue
        if nxt <= today:
            out.append(row)
    return out


def cmd_review(args) -> None:
    rows = due_mistakes()
    if not rows:
        print("[复习] 今日没有到期错题 —— 保持清零节奏，可以去推进新内容。")
        return
    print(f"\n=== 今日到期复习 {len(rows)} 条 ===\n")
    for i, row in enumerate(rows, 1):
        print(f"{i:>2}. {row['错题ID']} | {row['科目']}·{row['模块']}·{row['考点']}")
        print(f"    错因：{row['根本原因']}")
        print(f"    口诀：{row['避坑口诀']}")
        print(f"    阶段 {row['复习阶段']}/{len(REVIEW_INTERVALS)}  来源：{row['来源']}\n")

    if args.done:
        count = min(args.done, len(rows))
        all_rows = read_rows("mistakes")
        targets = {r["错题ID"] for r in rows[:count]}
        for row in all_rows:
            if row["错题ID"] not in targets:
                continue
            stage = to_int(row["复习阶段"]) + 1
            row["复习阶段"] = str(stage)
            if stage >= len(REVIEW_INTERVALS):
                row["已清零"] = "是"
                row["下次复习"] = ""
            else:
                gap = REVIEW_INTERVALS[stage]
                row["下次复习"] = (date.today() + timedelta(days=gap)).isoformat()
        replace_rows("mistakes", all_rows)
        print(f"[已更新] 标记 {count} 条完成复习")


def cmd_list_mistakes(args) -> None:
    rows = read_rows("mistakes")
    if args.subject:
        rows = [r for r in rows if r["科目"] == args.subject]
    if not rows:
        print("[错题本] 暂无数据")
        return
    print(f"\n{len(rows)} 条错题：\n")
    for row in rows[-args.limit:]:
        flag = "✔" if row.get("已清零") == "是" else " "
        print(f"[{flag}] {row['错题ID']} {row['录入日期']} {row['科目']}·{row['模块']} "
              f"| {row['考点']} | {row['错误类型']}")


# --------------------------------------------------------------------------- #
# 学习时长 / 模考
# --------------------------------------------------------------------------- #
def cmd_add_session(args) -> None:
    if args.yes:
        row = {
            "日期": args.date or date.today().isoformat(),
            "科目": args.subject, "内容": args.activity,
            "实际时长(分)": args.minutes, "计划时长(分)": args.planned or args.minutes,
            "专注度(1-5)": args.focus or 4, "备注": args.note or "-",
        }
    else:
        print("\n=== 录入学习时长 ===")
        subject = ask("科目", choices=list(MODULES.keys()) + ["综合"])
        activity = ask("内容（如：数学几何专项 / 英语二 2018 Text2 精读）")
        minutes = ask("实际时长(分钟)")
        planned = ask("计划时长(分钟)", default=minutes)
        focus = ask("专注度 1-5", default="4")
        note = ask("备注", default="-")
        row = {
            "日期": date.today().isoformat(), "科目": subject, "内容": activity,
            "实际时长(分)": minutes, "计划时长(分)": planned,
            "专注度(1-5)": focus, "备注": note,
        }
    append_row("sessions", row)
    print(f"[已录入] {row['日期']} {row['科目']} {row['实际时长(分)']}min")


def cmd_add_paper(args) -> None:
    if args.yes:
        row = {
            "日期": args.date or date.today().isoformat(),
            "科目": args.subject, "试卷名称": args.paper,
            "用时(分)": args.minutes, "得分": args.score, "满分": args.total,
            "错题数": args.wrong, "备注": args.note or "-",
        }
    else:
        print("\n=== 录入模考成绩 ===")
        subject = ask("科目", choices=["数学", "逻辑", "写作", "管综", "英语二", "六级"])
        paper = ask("试卷名称（如：2021 管综数学真题）")
        minutes = ask("用时(分钟)")
        score = ask("得分")
        total = ask("满分", default="75" if subject == "数学" else "100")
        wrong = ask("错题数")
        note = ask("备注", default="-")
        row = {
            "日期": date.today().isoformat(), "科目": subject, "试卷名称": paper,
            "用时(分)": minutes, "得分": score, "满分": total,
            "错题数": wrong, "备注": note,
        }
    append_row("papers", row)
    print(f"[已录入] {row['日期']} {row['科目']} {row['试卷名称']} "
          f"{row['得分']}/{row['满分']} 错 {row['错题数']} 题")


# --------------------------------------------------------------------------- #
# 报表
# --------------------------------------------------------------------------- #
def countdown_lines(cfg: dict) -> list[str]:
    today = date.today()
    lines = []
    for label, key in (("考研初试", "exam_kaoyan"),
                       ("六级笔试", "exam_cet6_written"),
                       ("六级口试", "exam_cet6_oral")):
        try:
            d = datetime.strptime(cfg[key], "%Y-%m-%d").date()
        except Exception:
            continue
        lines.append(f"  {label:<8} {d.isoformat()}   剩 {max((d - today).days, 0)} 天")
    return lines


def collect_stats() -> dict:
    mistakes = read_rows("mistakes")
    sessions = read_rows("sessions")
    papers = read_rows("papers")
    coverage = read_rows("coverage")

    by_module = Counter(f"{r['科目']}·{r['模块']}" for r in mistakes)
    by_etype = Counter(r.get("错误类型", "") for r in mistakes)
    by_point = Counter(f"{r['科目']}·{r['模块']}·{r['考点']}" for r in mistakes)

    today = date.today()
    week_start = today - timedelta(days=today.weekday())
    week_min = sum(to_int(r["实际时长(分)"]) for r in sessions
                   if parse_date(r["日期"]) >= week_start)
    total_min = sum(to_int(r["实际时长(分)"]) for r in sessions)
    by_subject_min = Counter()
    for r in sessions:
        by_subject_min[r["科目"]] += to_int(r["实际时长(分)"])

    latest: dict[str, dict] = {}
    for r in papers:
        latest[r["科目"]] = r  # 后写的覆盖先写的

    return {
        "mistakes": mistakes, "sessions": sessions, "papers": papers,
        "coverage": coverage, "by_module": by_module, "by_etype": by_etype,
        "by_point": by_point, "week_min": week_min, "total_min": total_min,
        "by_subject_min": by_subject_min, "latest": latest,
        "due": due_mistakes(today),
    }


def cmd_report(args) -> None:
    cfg = load_config()
    stats = collect_stats()
    w = 66
    print("=" * w)
    print("  2027 考研备考仪表盘".center(w - 4))
    print("=" * w)

    print("\n【倒计时】")
    for line in countdown_lines(cfg):
        print(line)

    print("\n【目标 vs 最新模考】")
    targets = cfg["targets"]
    for subject in ("数学", "逻辑", "写作", "管综", "管综总分", "英语二", "六级"):
        if subject not in targets:
            continue
        row = stats["latest"].get(subject)
        if row:
            got, total = to_float(row["得分"]), to_float(row["满分"], 100)
            gap = to_float(targets[subject]) - got
            mark = "✔" if gap <= 0 else "缺"
            print(f"  {subject:<8} 目标 {targets[subject]:>5}  |  最新 {got:>5.0f}/{total:.0f}"
                  f"  |  {mark} {gap:+.0f}   ({row['日期']})")
        else:
            print(f"  {subject:<8} 目标 {targets[subject]:>5}  |  尚无模考记录")

    print("\n【错题总览】")
    total_m = len(stats["mistakes"])
    cleared = sum(1 for r in stats["mistakes"] if r.get("已清零") == "是")
    print(f"  累计错题 {total_m} 条 | 已清零 {cleared} 条 | 今日到期复习 {len(stats['due'])} 条")
    if stats["by_module"]:
        print("  最薄弱模块 TOP5：")
        for name, cnt in stats["by_module"].most_common(5):
            print(f"    {name:<24} {cnt} 次")
    if stats["by_etype"]:
        print("  错误类型分布：")
        for name, cnt in stats["by_etype"].most_common():
            print(f"    {name:<12} {cnt} 次")

    print("\n【时长统计】")
    print(f"  本周累计 {stats['week_min'] // 60}h{stats['week_min'] % 60:02d}m"
          f" | 日均 {stats['week_min'] / 7 / 60:.1f}h"
          f" | 历史累计 {stats['total_min'] / 60:.1f}h")
    if stats["by_subject_min"]:
        print("  各科占比：")
        for name, mins in stats["by_subject_min"].most_common():
            pct = mins / max(stats["total_min"], 1) * 100
            print(f"    {name:<8} {mins / 60:>6.1f}h  {pct:>5.1f}%  {'█' * int(pct / 3)}")

    print("\n【周节奏提醒】")
    for k, v in STRATEGY.items():
        print(f"  {k}：{v}")
    print("\n" + "=" * w)


# --------------------------------------------------------------------------- #
# 时间表
# --------------------------------------------------------------------------- #
def schedule_markdown(key: str | None = None) -> str:
    keys = [key] if key else ["A", "B", "sunday"]
    lines = ["# 每日精确到分钟时间表", "",
             "> 由 `tools/tracker.py plan --export` 生成，请勿手工编辑；"
             "改时间请改脚本中的 `SCHEDULE`。", ""]
    for k in keys:
        block = SCHEDULE[k]
        lines += [f"## {block['title']}", "",
                  f"**{block['total']}**", "",
                  "| 时间 | 内容 | 执行要点 | 优先级 |",
                  "|---|---|---|---|"]
        for t, what, how, prio in block["rows"]:
            lines.append(f"| {t} | {what} | {how} | {prio} |")
        lines.append("")
    lines += ["## 执行规则", "",
              "1. **时间表是容器，不是枷锁**：被打断后从下一个整点块接上，禁止「反正乱了今天算了」。",
              "2. **优先级高于时长**：核心块（数学/逻辑/英语精读/写作）一天至少完成 3 个，才算达标日。",
              "3. **每晚 21:10 必须录入数据**：不录入 = 没发生。攒到周末补录的信息量会失真 70%。",
              "4. **连续 2 天未达标 → 立刻下调任务量**，而不是硬撑到崩盘。可持续 > 强度。", ""]
    return "\n".join(lines)


def cmd_plan(args) -> None:
    text = schedule_markdown(args.day)
    if args.export:
        ensure_dirs()
        out = PLANS / "每日时间表.md"
        out.write_text(text, encoding="utf-8")
        print(f"[已导出] {out}")
    else:
        print(text)


def cmd_today(args) -> None:
    today = date.today()
    key = args.day
    if key is None:
        wd = today.weekday()  # 0=周一 … 6=周日
        key = "sunday" if wd == 6 else ("A" if wd == 5 else "B")
    block = SCHEDULE[key]
    print("=" * 72)
    print(f"  {today.isoformat()} {['周一','周二','周三','周四','周五','周六','周日'][today.weekday()]}")
    print(f"  {block['title']}   （{block['total']}）")
    print("=" * 72)
    for t, what, how, prio in block["rows"]:
        print(f"  {t}  [{prio}] {what}")
        print(f"             ↳ {how}")
    due = due_mistakes(today)
    if due:
        print(f"\n  ⚠ 今日有 {len(due)} 条错题到期复习："
              f"{'、'.join(r['错题ID'] for r in due[:8])}")
    print()


# --------------------------------------------------------------------------- #
# Excel 仪表盘
# --------------------------------------------------------------------------- #
def cmd_excel(args) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[错误] 缺少 openpyxl，请先执行： python -m pip install openpyxl")
        return

    cfg = load_config()
    stats = collect_stats()
    ensure_dirs()
    out = REPORT / "备考仪表盘.xlsx"

    NAVY = "1F4E79"
    LIGHT = "EAF1F8"
    head_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="宋体", size=10)
    nav_font = Font(name="微软雅黑", size=13, bold=True, color=NAVY)
    thin = Side(style="thin", color="9DC3E6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor=NAVY)
    alt = PatternFill("solid", fgColor=LIGHT)

    wb = Workbook()

    # ---------- Sheet 1: 仪表盘 ----------
    ws = wb.active
    ws.title = "仪表盘"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col, width in zip("BCDEF", (18, 14, 14, 14, 34)):
        ws.column_dimensions[col].width = width

    r = 2
    ws[f"B{r}"] = "2027 考研备考仪表盘"
    ws[f"B{r}"].font = Font(name="微软雅黑", size=16, bold=True, color=NAVY)
    r += 1
    ws[f"B{r}"] = f"更新于 {date.today().isoformat()}"
    ws[f"B{r}"].font = Font(name="宋体", size=9, color="808080")
    r += 2

    def section(title: str) -> None:
        nonlocal r
        ws[f"B{r}"] = title
        ws[f"B{r}"].font = nav_font
        for col in "BCDEF":
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="DCE6F1")
            ws[f"{col}{r}"].border = border
        r += 1

    def table(headers: list[str], rows: list[list], start_col: int = 2) -> None:
        nonlocal r
        for i, h in enumerate(headers):
            c = ws.cell(row=r, column=start_col + i, value=h)
            c.font = head_font
            c.fill = fill
            c.border = border
            c.alignment = Alignment(horizontal="center", vertical="center")
        r += 1
        for j, row in enumerate(rows):
            for i, v in enumerate(row):
                c = ws.cell(row=r, column=start_col + i, value=v)
                c.font = body_font
                c.border = border
                if j % 2 == 1:
                    c.fill = alt
                c.alignment = Alignment(horizontal="left", vertical="center",
                                        wrap_text=False)
            r += 1
        r += 1

    section("一、倒计时")
    cd = []
    for label, key in (("考研初试", "exam_kaoyan"), ("六级笔试", "exam_cet6_written"),
                       ("六级口试", "exam_cet6_oral")):
        try:
            d = datetime.strptime(cfg[key], "%Y-%m-%d").date()
        except Exception:
            continue
        cd.append([label, d.isoformat(), max((d - date.today()).days, 0)])
    table(["节点", "日期", "剩余天数"], cd)

    section("二、目标达成度")
    rows = []
    for subject, target in cfg["targets"].items():
        rec = stats["latest"].get(subject)
        if rec:
            got = to_float(rec["得分"])
            rows.append([subject, target, round(got, 1), round(target - got, 1), rec["日期"]])
        else:
            rows.append([subject, target, "—", "—", "尚无记录"])
    table(["科目", "目标分", "最新得分", "差距", "最近模考"], rows)

    section("三、错题总览")
    total_m = len(stats["mistakes"])
    cleared = sum(1 for x in stats["mistakes"] if x.get("已清零") == "是")
    table(["指标", "数值"], [
        ["累计错题", total_m],
        ["已清零", cleared],
        ["待清零", total_m - cleared],
        ["今日到期复习", len(stats["due"])],
    ])

    section("四、薄弱模块 TOP")
    table(["科目·模块", "错题数"], [[k, v] for k, v in stats["by_module"].most_common(10)])
    mod_anchor = r - len(stats["by_module"].most_common(10)) - 2

    section("五、错误类型分布")
    table(["错误类型", "次数"], [[k, v] for k, v in stats["by_etype"].most_common()])

    section("六、时长统计")
    table(["指标", "数值"], [
        ["本周累计(小时)", round(stats["week_min"] / 60, 1)],
        ["本周日均(小时)", round(stats["week_min"] / 7 / 60, 1)],
        ["历史累计(小时)", round(stats["total_min"] / 60, 1)],
    ])
    table(["科目", "累计小时"], [[k, round(v / 60, 1)]
                              for k, v in stats["by_subject_min"].most_common()])

    # 柱状图：各模块错题数
    n_mod = len(stats["by_module"].most_common(10))
    if n_mod:
        chart = BarChart()
        chart.title = "各模块错题数"
        chart.y_axis.title = "错题数"
        chart.height, chart.width = 8, 16
        data = Reference(ws, min_col=3, min_row=mod_anchor,
                         max_row=mod_anchor + n_mod)
        cats = Reference(ws, min_col=2, min_row=mod_anchor + 1,
                         max_row=mod_anchor + n_mod)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"H{mod_anchor - 1}")

    # ---------- 其余 Sheet：全量数据 ----------
    col_map = {
        "错题本": "mistakes", "学习时长": "sessions",
        "模考记录": "papers", "知识点覆盖": "coverage",
    }
    for sheet_name, key in col_map.items():
        sh = wb.create_sheet(sheet_name)
        headers = HEADERS[key]
        for i, h in enumerate(headers, 1):
            c = sh.cell(row=1, column=i, value=h)
            c.font = head_font
            c.fill = fill
            c.border = border
            c.alignment = Alignment(horizontal="center")
            sh.column_dimensions[get_column_letter(i)].width = \
                max(10, min(30, max(len(h) * 2 + 4, 12)))
        for j, row in enumerate(read_rows(key), start=2):
            for i, h in enumerate(headers, 1):
                c = sh.cell(row=j, column=i, value=row.get(h, ""))
                c.font = body_font
                c.border = border
        sh.freeze_panes = "A2"
        if sh.max_row > 1:
            sh.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sh.max_row}"

    wb.save(out)
    print(f"[已生成] {out}")
    print("         包含 5 个工作表：仪表盘 / 错题本 / 学习时长 / 模考记录 / 知识点覆盖")


# --------------------------------------------------------------------------- #
# 演示数据（用于验证工具链，可随时用 init --force 清空）
# --------------------------------------------------------------------------- #
DEMO = {
    "mistakes": [
        ["E0001", "2026-09-17", "数学", "几何", "球内接正方体大圆截面", "2021 真题 15", "方法缺失",
         "在立体图中硬想，没有画二维截面降维", "球的问题先画大圆截面，把三维压成三角形", "0", "2026-09-18", "否", "示例"],
        ["E0002", "2026-09-17", "数学", "数据分析", "分组均分计数", "2020 真题 11", "计算失误",
         "均分问题忘记除以 A_k^k", "见「均匀分组」先问一句：组数要不要除阶乘", "0", "2026-09-18", "否", "示例"],
        ["E0003", "2026-09-17", "数学", "条件充分性判断", "均值不等式等号条件", "2019 真题 17", "陷阱误判",
         "用均值不等式没检验等号成立条件", "均值不等式出手前，先写等号成立条件", "0", "2026-09-18", "否", "示例"],
        ["E0004", "2026-09-17", "逻辑", "形式逻辑", "假言命题逆否等价", "2018 真题 28", "概念不清",
         "把「A→B」当成「B→A」用", "箭头只有反向加否定才等价：A→B ≡ ¬B→¬A", "0", "2026-09-18", "否", "示例"],
        ["E0005", "2026-09-17", "逻辑", "论证逻辑", "绝对量 vs 相对量", "2017 真题 33", "陷阱误判",
         "用绝对数量增长推出比例上升", "看到「比例/占比」立刻查基数有没有变", "0", "2026-09-18", "否", "示例"],
        ["E0006", "2026-09-17", "英语二", "阅读", "转折处出题点", "2019 Text2", "审题偏差",
         "只读了转折前的句子就选答案", "but / however 之后必有考点，先划出来", "0", "2026-09-18", "否", "示例"],
    ],
    "sessions": [
        ["2026-09-15", "数学", "几何专项：截面与割补", "90", "90", "4", "示例"],
        ["2026-09-15", "逻辑", "形式逻辑：假言逆否", "90", "90", "4", "示例"],
        ["2026-09-15", "英语二", "2019 Text2 精读", "100", "100", "3", "示例"],
        ["2026-09-16", "数学", "概率与排列组合", "90", "90", "4", "示例"],
        ["2026-09-16", "六级", "听力精听 1 套", "60", "60", "3", "示例"],
        ["2026-09-16", "写作", "论证有效性分析 1 篇", "90", "90", "4", "示例"],
        ["2026-09-17", "数学", "条件充分性判断专项", "90", "90", "5", "示例"],
        ["2026-09-17", "英语二", "2020 Text1 精读", "100", "100", "4", "示例"],
    ],
    "papers": [
        ["2026-09-13", "数学", "2021 管综数学真题", "58", "57", "75", "6", "示例"],
        ["2026-09-13", "逻辑", "2021 管综逻辑真题", "55", "46", "60", "7", "示例"],
        ["2026-09-14", "英语二", "2019 英语二真题", "170", "71", "100", "8", "示例"],
        ["2026-09-14", "六级", "2025.12 六级真题", "130", "468", "710", "—", "示例"],
    ],
}


def cmd_seed_demo(args) -> None:
    init_files()
    for key, rows in DEMO.items():
        replace_rows(key, [dict(zip(HEADERS[key], row)) for row in rows])
        print(f"[写入示例] {key}.csv  {len(rows)} 行")
    print("\n提示：验证完用 `python tools/tracker.py init --force` 可清空重来。")


def cmd_seed_baseline(args) -> None:
    """把 2017–2021 真题错点分析导入知识点覆盖表（掌握度标记为 0 = 未测）。"""
    init_files()
    rows = read_rows("coverage")
    existing = {(r["科目"], r["模块"], r["考点"]) for r in rows}
    added = 0
    for subject, module, point in BASELINE_COVERAGE:
        if (subject, module, point) in existing:
            continue
        rows.append({
            "科目": subject, "模块": module, "考点": point,
            "掌握度(1-5)": "0", "最近练习": "", "累计错次": "1",
            "备注": "2017–2021 真题错点反推（掌握度待摸底校准）",
        })
        added += 1
    replace_rows("coverage", rows)
    total = len([r for r in rows if r["科目"] in ("数学", "逻辑")])
    print(f"[基线导入] 新增 {added} 个已知错点，覆盖表共 {total} 个数学/逻辑考点")
    print("           掌握度全部为 0（未测）—— 摸底后请改成分数：")
    print("           5=稳定全对  4=偶有失误  3=需看提示  2=常错  1=基本不会  0=未测")
    print("\n下一步：按 plans/摸底测方案.md 做 4 天模块级摸底，再回填真实正确率。")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="tracker.py",
        description="2027 考研（管综 199 + 英语二 + CET-6）备考数据中枢",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("init", help="初始化 data/*.csv 与 config.json")
    s.add_argument("--force", action="store_true", help="覆盖已有文件")
    s.set_defaults(func=lambda a: init_files(a.force))

    s = sub.add_parser("today", help="打印今日分钟级时间表")
    s.add_argument("-d", "--day", choices=["A", "B", "sunday"], default=None)
    s.set_defaults(func=cmd_today)

    s = sub.add_parser("plan", help="打印/导出完整时间表")
    s.add_argument("-d", "--day", choices=["A", "B", "sunday"], default=None)
    s.add_argument("--export", action="store_true")
    s.set_defaults(func=cmd_plan)

    s = sub.add_parser("add-mistake", help="录入错题")
    s.add_argument("--subject"), s.add_argument("--module")
    s.add_argument("--point"), s.add_argument("--source")
    s.add_argument("--error-type", dest="error_type", default="概念不清")
    s.add_argument("--cause"), s.add_argument("--fix"), s.add_argument("--note")
    s.add_argument("-y", "--yes", action="store_true", help="不做交互，用参数值")
    s.set_defaults(func=cmd_add_mistake)

    s = sub.add_parser("list-mistakes", help="列出错题")
    s.add_argument("--subject"), s.add_argument("--limit", type=int, default=20)
    s.set_defaults(func=cmd_list_mistakes)

    s = sub.add_parser("review", help="今日到期复习（艾宾浩斯）")
    s.add_argument("--done", type=int, default=0, help="把前 N 条标记为已复习")
    s.set_defaults(func=cmd_review)

    s = sub.add_parser("add-session", help="录入学习时长")
    s.add_argument("--subject"), s.add_argument("--activity")
    s.add_argument("--minutes"), s.add_argument("--planned")
    s.add_argument("--focus"), s.add_argument("--date"), s.add_argument("--note")
    s.add_argument("-y", "--yes", action="store_true")
    s.set_defaults(func=cmd_add_session)

    s = sub.add_parser("add-paper", help="录入模考成绩")
    s.add_argument("--subject"), s.add_argument("--paper")
    s.add_argument("--minutes"), s.add_argument("--score")
    s.add_argument("--total"), s.add_argument("--wrong")
    s.add_argument("--date"), s.add_argument("--note")
    s.add_argument("-y", "--yes", action="store_true")
    s.set_defaults(func=cmd_add_paper)

    s = sub.add_parser("report", help="控制台报表")
    s.set_defaults(func=cmd_report)

    s = sub.add_parser("excel", help="生成 Excel 仪表盘")
    s.set_defaults(func=cmd_excel)

    s = sub.add_parser("seed-demo", help="写入演示数据（验证用）")
    s.set_defaults(func=cmd_seed_demo)

    s = sub.add_parser("seed-baseline", help="导入 2017-2021 真题错点基线到知识点覆盖表")
    s.set_defaults(func=cmd_seed_baseline)

    return p


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    ensure_dirs()
    args.func(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
